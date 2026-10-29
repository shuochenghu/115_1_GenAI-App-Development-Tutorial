"""第9週教師驗證：本機文件、Notebook與AppTest；不發送真實API請求。"""

import ast
from contextlib import redirect_stdout
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from types import SimpleNamespace, ModuleType
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
WEEK = ROOT / "week09"
PACKAGE = WEEK / "week09_document_simple"
TEMP = ROOT / "output/week09_revision"
TEMP.mkdir(parents=True, exist_ok=True)
(TEMP / ".gitignore").write_text("*\n", encoding="utf-8")
# AppTest使用可寫的工作區暫存目錄；不啟動伺服器或真人瀏覽器。
tempfile.tempdir = str(TEMP.resolve())
from streamlit.testing.v1 import AppTest


def read_book(teacher):
    ending = "教師版" if teacher else "學生版"
    return json.loads(next(WEEK.glob(f"*課堂簡化版_{ending}.ipynb")).read_text(encoding="utf-8"))


def app_functions(teacher):
    source = (PACKAGE / ("teacher" if teacher else "student") / "app.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    # 只執行匯入與三個函式；UI由AppTest另測。
    nodes = [n for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef))]
    namespace = {}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "app-functions", "exec"), namespace)
    return namespace


def expect_error(call, error=ValueError):
    try:
        call()
    except error:
        return
    raise AssertionError(f"未拒絕輸入，應為{error.__name__}")


def check_notebooks():
    student, teacher = read_book(False), read_book(True)
    assert [c["id"] for c in student["cells"]] == [c["id"] for c in teacher["cells"]]
    for s_cell, t_cell in zip(student["cells"], teacher["cells"]):
        if s_cell["id"] not in ["w09s-cover", "w09s-a", "w09s-b"]:
            assert s_cell == t_cell, s_cell["id"]
    for book, is_teacher in [(student, False), (teacher, True)]:
        cells = book["cells"]
        assert len({c["id"] for c in cells}) == len(cells)
        by_tag = {c["id"].removeprefix("w09s-"): c for c in cells}
        sequential = {}
        with tempfile.TemporaryDirectory(dir=TEMP) as folder, patch.dict(os.environ, {"OPENAI_API_KEY": ""}):
            old = Path.cwd()
            os.chdir(folder)
            try:
                for cell in cells:
                    source = "".join(cell["source"])
                    assert "\ufffd" not in source and "???" not in source
                    if cell["cell_type"] == "markdown":
                        headings = re.findall(r"^#{1,6} .+", source, re.M)
                        if headings:
                            assert len(headings) == 1 and source.startswith("#"), cell["id"]
                        continue
                    assert cell["outputs"] == [] and cell["execution_count"] is None
                    ast.parse(source)
                    with redirect_stdout(StringIO()):
                        exec(compile(source, cell["id"], "exec"), sequential)
                    # 按宣告先備在乾淨namespace重跑，防止隱藏跨格依賴。
                    fresh = {}
                    for dependency in cell["metadata"]["teaching"]["requires"]:
                        exec("".join(by_tag[dependency]["source"]), fresh)
                    with redirect_stdout(StringIO()):
                        exec(source, fresh)
                assert sequential["summary_result"] is None
                assert sequential["summary_error"] is None
                assert (sequential["document_result"] is not None) == is_teacher
            finally:
                os.chdir(old)
        # Markdown閱讀版逐格保留原碼，App的A/B與Notebook相同。
        ending = "教師版" if is_teacher else "學生版"
        reading = next(WEEK.glob(f"*課堂簡化版_{ending}.md")).read_text(encoding="utf-8")
        app_tree = ast.parse((PACKAGE / ("teacher" if is_teacher else "student") / "app.py").read_text(encoding="utf-8"))
        for cell in cells:
            assert "".join(cell["source"]) in reading
        for tag, function in [("a", "clean_text"), ("b", "chunk_text")]:
            book_tree = ast.parse("".join(by_tag[tag]["source"]))
            expected = next(n for n in book_tree.body if isinstance(n, ast.FunctionDef))
            actual = next(n for n in app_tree.body if isinstance(n, ast.FunctionDef) and n.name == function)
            assert ast.dump(expected) == ast.dump(actual)
        assert ("raise NotImplementedError" in "".join(by_tag["a"]["source"])) != is_teacher
        assert ("raise NotImplementedError" in "".join(by_tag["b"]["source"])) != is_teacher
    print(f"Notebook：雙版各{len(teacher['cells'])}格，JSON／ID／輸出／AST／獨立先備／順序／閱讀版／答案分離通過")


def fixtures():
    """只建立記憶體中的reader測試資料，不作為可交付文件或掃描OCR成果。"""
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    from docx import Document

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Club meeting 4 people) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.add_blank_page(width=300, height=300)
    pdf_io = BytesIO()
    writer.write(pdf_io)

    blank_writer = PdfWriter()
    blank_writer.add_blank_page(width=300, height=300)
    blank_io = BytesIO()
    blank_writer.write(blank_io)
    encrypted_io = BytesIO()
    blank_writer.encrypt("classroom-only")
    blank_writer.write(encrypted_io)
    long_writer = PdfWriter()
    for _ in range(51):
        long_writer.add_blank_page(width=100, height=100)
    long_io = BytesIO()
    long_writer.write(long_io)

    document = Document()
    document.add_paragraph("段落前：活動日期10/28")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "人數"
    table.cell(0, 1).text = "4"
    document.add_paragraph("段落後：報名未公告")
    docx_io = BytesIO()
    document.save(docx_io)
    return {
        "pdf": pdf_io.getvalue(), "blank": blank_io.getvalue(),
        "encrypted": encrypted_io.getvalue(), "long": long_io.getvalue(), "docx": docx_io.getvalue(),
    }


def check_functions(data):
    namespace = app_functions(True)
    clean, chunk, extract = [namespace[x] for x in ["clean_text", "chunk_text", "extract_text"]]
    assert clean("A  B\r\n\r\n\r\nC\t4。") == "A B\n\nC 4。"
    # 不同長度、大小與重疊，驗證尾段、前進量與完整覆蓋。
    for text in ["", " ", "A", "ABCDEFGHIJ", "中文日期10/28\n\n每組4人。", " A B C "]:
        for size in range(1, 8):
            for overlap in range(size):
                result = chunk(text, size, overlap)
                if not text.strip():
                    assert result == []
                    continue
                covered = set()
                for number, item in enumerate(result):
                    assert item["chunk_id"] == number
                    assert item["text"] == text[item["start"]:item["end"]]
                    assert 0 < len(item["text"]) <= size
                    covered.update(range(item["start"], item["end"]))
                assert all(i in covered for i, char in enumerate(text) if not char.isspace())
    for size, overlap in [(0, 0), (-1, 0), (4, -1), (4, 4), (4, 5)]:
        expect_error(lambda: chunk("A", size, overlap))
    for name in ["note.TXT", "note.md"]:
        assert extract(name, "中文4人".encode("utf-8-sig")) == "中文4人"
        assert extract(name, "中文4人".encode("cp950"), "cp950") == "中文4人"
    expect_error(lambda: extract("x.txt", "社團".encode("cp950")), UnicodeDecodeError)
    assert "地點=圖書館,討論室" in extract("x.csv", '活動,地點\n交流,"圖書館,討論室"\n'.encode())
    for csv_text in [",地點\n活動,A\n", "活動,活動\nA,B\n", "活動,地點\nA\n", "活動\nA,B\n"]:
        expect_error(lambda: extract("x.csv", csv_text.encode()))
    expect_error(lambda: extract("x.csv", ("活動\n" + "A\n" * 201).encode()))
    expect_error(lambda: extract("x.exe", b"A"))
    pdf_text = extract("fixture.pdf", data["pdf"])
    assert "Club meeting 4 people" in pdf_text and "[第2頁]" in pdf_text and "此頁無文字層" in pdf_text
    for name in ["blank", "encrypted", "long"]:
        expect_error(lambda: extract("fixture.pdf", data[name]))
    word_text = extract("fixture.docx", data["docx"])
    assert word_text.index("段落前") < word_text.index("人數 | 4") < word_text.index("段落後")
    expect_error(lambda: extract("broken.pdf", b"invalid"), Exception)
    expect_error(lambda: extract("broken.docx", b"invalid"), Exception)
    for name in ["clean_text", "chunk_text"]:
        expect_error(lambda: app_functions(False)[name]("學生未完成"), NotImplementedError)
    with tempfile.TemporaryDirectory(dir=TEMP) as folder:
        directory = Path(folder)
        (directory / "sample.pdf").write_bytes(data["pdf"])
        (directory / "sample.docx").write_bytes(data["docx"])
        old = Path.cwd()
        os.chdir(folder)
        try:
            for tag in ["pdf-demo", "docx-demo"]:
                source = "".join(next(c for c in read_book(True)["cells"] if c["id"] == "w09s-" + tag)["source"])
                captured = StringIO()
                with redirect_stdout(captured):
                    exec(source, {})
                assert "未執行真實" not in captured.getvalue()
        finally:
            os.chdir(old)
    print("函式：清理、切塊邊界、UTF-8／CP950、CSV欄數／引號／上限、PDF文字／空白／加密／頁數、DOCX順序與損壞檔通過")


def button(at, label):
    return next(item for item in at.button if item.label == label)


def check_apps(data):
    import streamlit as st
    # download_button與上傳無法用AppTest模擬真人檔案對話框，攔截下載資料與替代上傳物件。
    downloads = []
    current = {"upload": None}

    def download(label, data, **kwargs):
        downloads.append(json.loads(data))
        return False

    def submit(at):
        downloads.clear()
        button(at, "處理文件").click().run()
        assert not at.exception, at.exception

    with patch.object(st, "download_button", side_effect=download), patch.object(st, "file_uploader", side_effect=lambda *a, **k: current["upload"]):
        teacher = AppTest.from_file(str(PACKAGE / "teacher/app.py"), default_timeout=20).run()
        assert not teacher.exception and not downloads
        submit(teacher)
        assert downloads and teacher.session_state["document_result"]["chunks"]
        first = downloads[-1]
        teacher.run()
        assert teacher.session_state["document_result"] == first
        teacher.number_input[0].set_value(0)
        submit(teacher)
        assert downloads[-1]["overlap"] == 0
        assert len(downloads[-1]["chunks"]) <= len(first["chunks"])
        teacher.number_input[0].set_value(120)
        submit(teacher)
        assert teacher.error and not downloads and teacher.session_state["document_result"] is None
        teacher.number_input[0].set_value(20)
        submit(teacher)
        assert downloads
        downloads.clear()
        button(teacher, "清除結果").click().run()
        assert teacher.session_state["document_result"] is None and not downloads

        teacher.checkbox[0].uncheck()
        for filename, content, expected in [
            ("x.csv", (PACKAGE / "sample_data/events.csv").read_bytes(), "活動=讀書交流"),
            ("x.pdf", data["pdf"], "Club meeting"),
            ("x.docx", data["docx"], "段落前"),
            ("x.txt", "自備中文活動4人".encode(), "4人"),
        ]:
            current["upload"] = SimpleNamespace(name=filename, getvalue=lambda content=content: content)
            submit(teacher)
            assert downloads and expected in downloads[-1]["cleaned_text"]

        for filename, content in [
            ("blank.txt", b"  \n"), ("long.txt", b"a" * 20001),
            ("large.txt", b"a" * (1024 * 1024 + 1)), ("blank.pdf", data["blank"]),
            ("broken.pdf", b"bad"), ("broken.docx", b"bad"), ("bad.csv", b"a,b\nx\n"),
            ("encoding.txt", "社團".encode("cp950")),
        ]:
            # 每次先成功，再提交錯誤，驗證沒有舊結果下載。
            current["upload"] = SimpleNamespace(name="ok.txt", getvalue=lambda: b"valid input")
            submit(teacher)
            assert downloads
            current["upload"] = SimpleNamespace(name=filename, getvalue=lambda content=content: content)
            submit(teacher)
            assert teacher.error and not downloads and teacher.session_state["document_result"] is None
        current["upload"] = SimpleNamespace(name="cp950.txt", getvalue=lambda: "社團4人".encode("cp950"))
        teacher.selectbox[0].select("cp950")
        submit(teacher)
        assert downloads[-1]["cleaned_text"] == "社團4人"
        current["upload"] = None
        submit(teacher)
        assert teacher.error and not downloads

        student = AppTest.from_file(str(PACKAGE / "student/app.py"), default_timeout=20).run()
        submit(student)
        assert student.warning and not downloads and student.session_state["document_result"] is None

        # 模擬學生把Notebook已完成A/B移入App，驗證轉移操作可直接成功。
        student_source = (PACKAGE / "student/app.py").read_text(encoding="utf-8")
        teacher_source = (PACKAGE / "teacher/app.py").read_text(encoding="utf-8")
        for name in ["clean_text", "chunk_text"]:
            tree = ast.parse(student_source)
            target = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == name)
            answer = next(node for node in ast.parse(teacher_source).body if isinstance(node, ast.FunctionDef) and node.name == name)
            lines = student_source.splitlines(keepends=True)
            student_source = "".join(lines[:target.lineno - 1]) + ast.get_source_segment(teacher_source, answer) + "\n" + "".join(lines[target.end_lineno:])
        completed = PACKAGE / "student/.verify_completed.py"
        try:
            completed.write_text(student_source, encoding="utf-8")
            completed_at = AppTest.from_file(str(completed), default_timeout=20).run()
            submit(completed_at)
            assert downloads and not completed_at.warning
        finally:
            completed.unlink(missing_ok=True)
    print("AppTest：雙版啟動、TODO阻擋、移入A/B成功、多格式替代上傳、下載來源、rerun保存、錯誤清舊結果及清除通過")


def check_api():
    source = "".join(next(c for c in read_book(True)["cells"] if c["id"] == "w09s-api")["source"])
    calls = []
    client_settings = []
    state = {"response": None}
    fake_module = ModuleType("openai")

    def create(**kwargs):
        calls.append(kwargs)
        response = state["response"]
        if isinstance(response, Exception):
            raise response
        return response

    def client(**kwargs):
        client_settings.append(kwargs)
        return SimpleNamespace(responses=SimpleNamespace(create=create))

    fake_module.OpenAI = client
    def run(response, paid=True, key="offline-placeholder", input_override=None):
        calls.clear()
        client_settings.clear()
        state["response"] = response
        activity = source.replace("RUN_PAID_API = False", f"RUN_PAID_API = {paid}")
        if input_override is not None:
            activity = activity.replace('summary_input = "社團下週三下午在圖書館討論室交流。報名方式尚未公告。"', "summary_input = " + repr(input_override))
        namespace = {"summary_result": "old-success"}
        with patch.dict(sys.modules, {"openai": fake_module}), patch("dotenv.load_dotenv"), patch.dict(os.environ, {"OPENAI_API_KEY": key}), redirect_stdout(StringIO()):
            exec(activity, namespace)
        return namespace

    good = SimpleNamespace(status="completed", output_text="三個重點", output=[])
    assert run(good, paid=False)["summary_result"] is None and not calls
    assert run(good)["summary_result"] == "三個重點" and len(calls) == 1
    assert calls[0]["store"] is False
    assert client_settings[0]["max_retries"] == 0 and client_settings[0]["timeout"] == 60.0
    for response in [
        SimpleNamespace(status="incomplete", output_text="部分", output=[]),
        SimpleNamespace(status="failed", output_text="錯誤", output=[]),
        SimpleNamespace(status="completed", output_text=" ", output=[]),
        SimpleNamespace(status="completed", output_text="文字", output=[SimpleNamespace(type="message", content=[SimpleNamespace(type="refusal")])]),
        RuntimeError("offline network error"),
    ]:
        result = run(response)
        assert result["summary_result"] is None and result["summary_error"] and len(calls) == 1
    assert run(good, key="")["summary_error"] and not calls
    for text in [" ", "A" * 6001]:
        assert run(good, input_override=text)["summary_error"] and not calls
    print("選做API：假SDK驗證預設0次、有效最多1次、參數、缺key／超長／拒答／未完成／空結果／失敗清舊資料通過")


def check_rebuild_and_links():
    targets = list(WEEK.glob("*課堂簡化版_*.ipynb")) + list(WEEK.glob("*課堂簡化版_*.md")) + list(PACKAGE.rglob("app.py")) + list((PACKAGE / "sample_data").iterdir())
    before = {path: path.read_bytes() for path in targets}
    subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "scripts/build_week09_simple.py")], check=True, capture_output=True)
    assert all(path.read_bytes() == content for path, content in before.items())
    for path in [*WEEK.glob("*課堂簡化版*.md"), PACKAGE / "README.md"]:
        text = path.read_text(encoding="utf-8")
        assert "\ufffd" not in text and "???" not in text
        assert len(re.findall(r"^```", text, re.M)) % 2 == 0
        for target in re.findall(r"\]\(([^)]+)\)", text):
            if not target.startswith(("https://", "http://", "#")):
                assert (path.parent / target.split("#")[0]).exists(), (path, target)
    snapshot = TEMP / "before.json"
    if snapshot.exists():
        existing = json.loads(snapshot.read_text(encoding="utf-8"))
        count = 0
        for path, checksum in existing.items():
            # 只鎖定本次應保留的第9週原檔；其他週可能在另一個聊天同時修改。
            if Path(path).parts[0] != "week09":
                continue
            assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == checksum, path
            count += 1
        print(f"原第9週：{count}份檔案SHA-256不變；重建與本機連結通過")


def main():
    check_notebooks()
    data = fixtures()
    check_functions(data)
    check_apps(data)
    check_api()
    check_rebuild_and_links()
    print("全部離線驗證通過；0次真實API，未啟動伺服器，未測真人瀏覽器／Colab／Cloud／試教。")


if __name__ == "__main__":
    main()
