"""第 7–12 週課堂環境的離線檢核；在課程 repo 根目錄執行。

使用同一個 Python 逐一開新程序，避免不同週次同名 helper 互相干擾。
測試不讀取真實金鑰、不呼叫付費 API；AppTest 不等於瀏覽器或雲端部署驗證。
"""

from __future__ import annotations

import argparse
from contextlib import ExitStack
from importlib.util import find_spec
from io import BytesIO
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
from unittest.mock import patch
from urllib.request import urlopen


ROOT = Path(__file__).resolve().parents[1]
APPS = (
    "week07/week07_streamlit_app",
    "week08/week08_midterm_starter",
    "week08/week08_midterm_example_summarizer",
    "week09/week09_document_processor",
    "week10/week10_semantic_search_app",
    "week11/week11_rag_app",
    "week12/week12_vision_app",
)


def check_readers() -> None:
    """用記憶體內的小型文件確認 reader 與目前套件版本相容。"""
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from document_utils import extract_text, clean_text, chunk_text

    assert "課堂" in extract_text("test.txt", "課堂測試".encode("cp950"))
    assert "課堂" in extract_text("test.md", "課堂測試".encode("utf-8"))
    for encoding in ("utf-8", "cp950"):
        assert "課堂" in extract_text("test.csv", "主題,週次\n課堂,9\n".encode(encoding))
    word = Document()
    word.add_paragraph("課堂文件")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "表格內容"
    buffer = BytesIO()
    word.save(buffer)
    assert "表格內容" in extract_text("test.docx", buffer.getvalue())

    # PDF fixture 只用內建字型與英文字，避免外部字型檔影響環境測試。
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Classroom PDF) Tj ET")
    page[NameObject("/Contents")] = stream
    buffer = BytesIO()
    pdf.write(buffer)
    assert "Classroom PDF" in extract_text("test.pdf", buffer.getvalue())
    blank = PdfWriter()
    blank.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    blank.write(buffer)
    try:
        extract_text("blank.pdf", buffer.getvalue())
    except ValueError as exc:
        assert "OCR" in str(exc)
    else:
        raise AssertionError("無文字 PDF 應提示 OCR")
    chunks = chunk_text(clean_text("課堂測試\n\n" * 30), chunk_size=40, overlap=8)
    assert len(chunks) > 1 and all(c["end"] > c["start"] for c in chunks)


def check_app(app_dir: Path) -> None:
    """跑初始畫面、上傳邊界及離線檢索，不碰 API 或使用者 secrets。"""
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    sys.path.insert(0, str(app_dir))
    with ExitStack() as stack:
        stack.enter_context(patch.object(st, "secrets", {}))
        stack.enter_context(patch("dotenv.load_dotenv", return_value=False))
        stack.enter_context(patch.dict(os.environ, {"OPENAI_API_KEY": "", "ANONYMIZED_TELEMETRY": "False"}))
        # 即使日後 App 意外在初始化呼叫 API，此檢核也會立即失敗，不產生費用。
        stack.enter_context(patch("openai.OpenAI", side_effect=AssertionError("離線測試不可建立 API client")))
        app = AppTest.from_file(str(app_dir / "app.py"), default_timeout=30).run()
        assert not app.exception, [e.message for e in app.exception]
        # 第 9 週側欄固定用 st.error 顯示安全提醒，不能誤判為執行失敗。
        baseline_errors = [e.value for e in app.error]
        if app_dir.parent.name not in {"week09", "week10", "week11", "week12"}:
            return
        assert app.file_uploader[0].proto.max_upload_size_mb == 8

        limit = 8 * 1024 * 1024
        if app_dir.parent.name == "week12":
            from PIL import Image

            buffer = BytesIO()
            Image.new("RGB", (320, 180), "white").save(buffer, format="PNG")
            image_bytes = buffer.getvalue()
            # PNG 結尾後的 padding 不影響解碼，可精準測試 byte 上限。
            boundary_image = image_bytes + b" " * (limit - len(image_bytes))
            app.file_uploader[0].set_value(
                ("boundary.png", boundary_image, "image/png")
            ).run()
            assert not app.exception and [e.value for e in app.error] == baseline_errors
            assert len(app.metric) >= 2
            app.file_uploader[0].set_value(
                ("too_large.png", boundary_image + b"x", "image/png")
            ).run()
            assert not app.exception
            assert any("8 MB" in e.value for e in app.error), "超限圖片應被拒絕"
            return

        check_readers()

        # 大量尾端空白會在前處理移除，讓邊界測試不產生數萬個 chunks。
        content = b"Classroom boundary" + b" " * (limit - len(b"Classroom boundary"))
        app.file_uploader[0].set_value(("boundary.txt", content, "text/plain")).run()
        assert not app.exception and [e.value for e in app.error] == baseline_errors, ([e.message for e in app.exception], [e.value for e in app.error])
        app.file_uploader[0].set_value(("too_large.txt", content + b"x", "text/plain")).run()
        assert not app.exception
        assert any("8 MB" in e.value for e in app.error), "超限檔案應被拒絕"

        sample = next((app_dir / "sample_data").glob("*"))
        app.file_uploader[0].set_value((sample.name, sample.read_bytes(), "text/plain")).run()
        assert not app.exception and [e.value for e in app.error] == baseline_errors, ([e.message for e in app.exception], [e.value for e in app.error])
        if app_dir.parent.name == "week09":
            assert len(app.metric) >= 2 and len(app.get("download_button")) >= 1
            return

        next(b for b in app.button if "索引" in b.label).click().run()
        assert not app.exception and [e.value for e in app.error] == baseline_errors, ([e.message for e in app.exception], [e.value for e in app.error])
        if app_dir.parent.name == "week10":
            indexed = app.session_state["indexed_chunks"]
            assert indexed
            app.text_input[0].set_value("第 11 週").run()
            next(b for b in app.button if b.label == "搜尋").click().run()
            assert not app.exception and len(app.expander) > 0
            if find_spec("chromadb") is None:
                print("SKIP ChromaDB 選讀：此環境只安裝第 10 週必要依賴。", flush=True)
                return
            from embedding_utils import build_chroma_collection
            collection = build_chroma_collection(indexed)
            result = collection.query(query_embeddings=[indexed[0]["embedding"]], n_results=1)
            assert abs(result["distances"][0][0]) < 1e-5
        else:
            from embedding_utils import query_chroma_index
            from rag_utils import build_rag_context, format_sources
            collection = app.session_state["rag_collection"]
            chunks = app.session_state["rag_chunks"]
            hits = query_chroma_index(collection, chunks[0]["text"], min_score=-1)
            assert hits and abs(hits[0]["distance"]) < 1e-5
            assert "[來源 1]" in build_rag_context(hits)
            assert format_sources(hits)


def check_cli_config(app_dir: Path) -> None:
    """經真正 CLI 的解析流程確認設定來源，攔住伺服器啟動以便讀出結果。"""
    from click.testing import CliRunner
    from streamlit import config
    from streamlit.web import cli

    expected = app_dir / ".streamlit" / "config.toml"

    def inspect_config(*args, **kwargs):
        config.get_config_options(force_reparse=True)
        if expected.exists():
            assert config.get_option("server.maxUploadSize") == 8
            assert Path(config.get_where_defined("server.maxUploadSize")).resolve() == expected

    entry = "app.py" if Path.cwd() == app_dir else (app_dir / "app.py").relative_to(ROOT).as_posix()
    with patch("streamlit.web.bootstrap.run", side_effect=inspect_config) as mocked:
        result = CliRunner().invoke(cli.main, ["run", entry, "--browser.gatherUsageStats=false"])
        assert result.exit_code == 0, str(result.exception)
        mocked.assert_called_once()


def check_server(app_dir: Path) -> None:
    """實際啟動本機伺服器並檢查健康端點；不把它視為瀏覽器互動測試。"""
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    entry = "app.py" if Path.cwd() == app_dir else (app_dir / "app.py").relative_to(ROOT).as_posix()
    args = [sys.executable, "-m", "streamlit", "run", entry, "--server.headless=true",
            "--browser.gatherUsageStats=false", "--server.address=127.0.0.1", f"--server.port={port}"]
    with tempfile.TemporaryFile() as log:
        proc = subprocess.Popen(args, stdout=log, stderr=log,
                                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        try:
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                if proc.poll() is not None:
                    break
                try:
                    with urlopen(f"http://127.0.0.1:{port}/_stcore/health", timeout=1) as response:
                        if response.status == 200:
                            return
                except OSError:
                    time.sleep(0.2)
            log.seek(0)
            raise AssertionError(log.read().decode("utf-8", errors="replace"))
        finally:
            proc.terminate()
            try:
                proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                proc.kill()
                proc.wait()


def main() -> None:
    # Windows 主控台也使用 UTF-8，避免測試紀錄中的繁體中文變成亂碼。
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", choices=APPS)
    parser.add_argument("--cwd", choices=("repo", "app"), default="repo")
    options = parser.parse_args()
    if options.worker:
        app_dir = ROOT / options.worker
        os.chdir(ROOT if options.cwd == "repo" else app_dir)
        check_cli_config(app_dir)
        check_app(app_dir)
        check_server(app_dir)
        print(f"PASS {options.worker} cwd={options.cwd}", flush=True)
        return
    subprocess.run([sys.executable, "-m", "pip", "check"], check=True)
    for app in APPS:
        for cwd in ("repo", "app"):
            subprocess.run([sys.executable, str(Path(__file__).resolve()), "--worker", app, "--cwd", cwd], check=True)
    print("PASS: 七個 App、兩種工作目錄的離線檢核完成；尚未驗證瀏覽器、Colab、付費 API 或雲端。")


if __name__ == "__main__":
    main()
